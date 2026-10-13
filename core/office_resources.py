"""Criacao de documentos e planilhas, sem macros, codigo ou sobrescrita implicita."""
import csv
import io
import json
import math
import os
import tempfile
from pathlib import Path
from core.file_library import safe_path
from core.resource_jobs import request
from core.document_reader import read


def write_new(path,callback):
    path=safe_path(path)
    if path.exists() or not path.parent.is_dir():raise ValueError('Destino deve ser novo e sua pasta deve existir.')
    handle,temporary=tempfile.mkstemp(suffix=path.suffix,dir=path.parent)
    os.close(handle)
    try:
        callback(Path(temporary))
        # Exclusive creation avoids replacing a file created during the operation.
        with open(temporary,'rb') as source,path.open('xb') as target:
            import shutil
            shutil.copyfileobj(source,target)
    finally:Path(temporary).unlink(missing_ok=True)
    return {'ok':True,'path':str(path),'bytes':path.stat().st_size,'verified_exists':True}


def run(p):
    action=p.get('action','read')
    if action=='status':
        import importlib.util
        return {'available':True,'dependencies':{name:bool(importlib.util.find_spec(name)) for name in ('docx','openpyxl','PIL')}}
    if action=='read':return read(p)
    if action=='compare':
        left=read({'path':str(safe_path(p.get('path'))),'max_chars':20000})
        right=read({'path':str(safe_path(p.get('other'))),'max_chars':20000})
        import difflib
        diff=list(difflib.unified_diff(left['text'].splitlines(),right['text'].splitlines(),fromfile=left['path'],tofile=right['path'],lineterm=''))
        return {'diff':'\n'.join(diff)[:20000],'complete':left['complete'] and right['complete'] and len('\n'.join(diff))<=20000,'data_is_untrusted':True}
    path=None if action=='statistics' else safe_path(p.get('path'))
    if action=='edit_spreadsheet':
        from core.document_reader import check_archive,MAX_BYTES
        from core.file_library import fingerprint
        import re
        source=safe_path(p.get('other'))
        if source.suffix.lower()!='.xlsx' or path.suffix.lower()!='.xlsx' or not source.is_file() or source.stat().st_size>MAX_BYTES:
            raise ValueError('Informe XLSX existente e destino XLSX novo; macros nao sao aceitas.')
        check_archive(source)
        edits=json.loads(p.get('edits_json','[]'))
        if not isinstance(edits,list) or not 1<=len(edits)<=500:raise ValueError('Informe de uma a quinhentas edicoes de celula.')
        from openpyxl.utils.cell import coordinate_to_tuple
        for item in edits:
            if not isinstance(item,dict) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,3}',str(item.get('cell',''))):raise ValueError('Celula invalida; use A1, B2 etc.')
            row,column=coordinate_to_tuple(item['cell'])
            if row>5000 or column>100:raise ValueError('Edicao limitada a 5000 linhas e 100 colunas.')
            value=item.get('value')
            if value is not None and (not isinstance(value,(str,int,float,bool)) or isinstance(value,float) and not math.isfinite(value) or isinstance(value,str) and len(value)>32000):raise ValueError('Valor de celula invalido.')
        initial=fingerprint(source)
        def save(target):
            if fingerprint(source)!=initial:raise ValueError('Planilha mudou depois de preparada a edicao; revise.')
            from openpyxl import load_workbook
            book=load_workbook(source,keep_links=False)
            try:
                name=p.get('sheet') or book.sheetnames[0]
                if name not in book.sheetnames:raise ValueError('Planilha interna nao encontrada; consulte o documento primeiro.')
                sheet=book[name]
                for item in edits:
                    cell=sheet[item['cell']];cell.value=item.get('value')
                    if isinstance(cell.value,str):cell.data_type='s'
                book.save(target)
            finally:book.close()
        return request('Editar copia da planilha',str(source)+' → '+str(path),lambda:write_new(path,save))
    if action=='document':
        if path.suffix.lower() not in ('.docx','.txt','.md'):raise ValueError('Escolha DOCX, TXT ou MD.')
        text=str(p.get('text',''))
        if not text or len(text)>100000:raise ValueError('Informe texto de ate 100 mil caracteres.')
        def save(target):
            if path.suffix.lower()=='.docx':
                from docx import Document
                doc=Document()
                if p.get('title'):doc.add_heading(str(p['title'])[:300],0)
                for line in text.splitlines():doc.add_paragraph(line)
                doc.save(target)
            else:target.write_text(text,encoding='utf-8')
        return request('Criar documento',str(path),lambda:write_new(path,save))
    if action in ('spreadsheet','statistics'):
        try:rows=json.loads(p.get('rows_json','[]'))
        except ValueError:raise ValueError('Envie uma matriz JSON de celulas.')
        if not isinstance(rows,list) or not 1<=len(rows)<=5000 or any(not isinstance(row,list) or len(row)>100 for row in rows):raise ValueError('Matriz acima dos limites ou formato invalido.')
        for row in rows:
            for cell in row:
                if cell is not None and not isinstance(cell,(str,int,float,bool)):raise ValueError('Celula deve conter texto, numero ou vazio.')
                if isinstance(cell,str) and len(cell)>32000:raise ValueError('Texto de celula muito longo.')
                if isinstance(cell,float) and not math.isfinite(cell):raise ValueError('Numero nao finito.')
        if action=='statistics':
            column=int(p.get('column',1))-1
            if not 0<=column<100:raise ValueError('Coluna invalida.')
            values=[row[column] for row in rows if column<len(row) and isinstance(row[column],(int,float)) and not isinstance(row[column],bool)]
            if not values:raise ValueError('Sem numeros na coluna escolhida.')
            return {'count':len(values),'sum':sum(values),'min':min(values),'max':max(values),'mean':sum(values)/len(values)}
        if path.suffix.lower() not in ('.xlsx','.csv'):raise ValueError('Escolha XLSX ou CSV.')
        def save(target):
            if path.suffix.lower()=='.xlsx':
                from openpyxl import Workbook
                book=Workbook();sheet=book.active;sheet.title='Dados'
                for row in rows:
                    sheet.append(row)
                    # Formula-looking text remains literal, not executable spreadsheet content.
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value,str):cell.data_type='s'
                sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
                if p.get('chart'):
                    if len(rows)<2 or len(rows[0])<2:raise ValueError('Grafico exige cabecalho e ao menos duas colunas.')
                    column=int(p.get('column',2))
                    if not 2<=column<=min(100,len(rows[0])) or any(column>len(row) or not isinstance(row[column-1],(int,float)) or isinstance(row[column-1],bool) for row in rows[1:]):
                        raise ValueError('A coluna do grafico deve conter somente numeros nos dados.')
                    from openpyxl.chart import BarChart,Reference
                    chart=BarChart();chart.title=str(p.get('title') or 'Dados')[:150];chart.y_axis.title='Valor';chart.x_axis.title='Categoria'
                    chart.add_data(Reference(sheet,min_col=column,max_col=column,min_row=1,max_row=len(rows)),titles_from_data=True)
                    chart.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=len(rows)));chart.y_axis.scaling.min=0
                    if any(row[column-1]<0 for row in rows[1:]):chart.y_axis.scaling.min=None
                    sheet.add_chart(chart,'D2')
                book.save(target);book.close()
            else:
                with target.open('w',newline='',encoding='utf-8-sig') as stream:
                    writer=csv.writer(stream)
                    for row in rows:writer.writerow(["'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')) else value for value in row])
        return request('Criar planilha',str(path),lambda:write_new(path,save))
    if action in ('convert_image','annotate_image'):
        source=safe_path(p.get('other'))
        if path.suffix.lower() not in ('.png','.jpg','.webp'):raise ValueError('Destino PNG, JPG ou WEBP.')
        marks=json.loads(p.get('annotations_json','[]')) if action=='annotate_image' else []
        if not isinstance(marks,list) or len(marks)>50:raise ValueError('Informe no maximo cinquenta anotacoes.')
        from core.file_library import fingerprint
        initial=fingerprint(source)
        def save(target):
            from PIL import Image,ImageOps,ImageDraw
            if fingerprint(source)!=initial:raise ValueError('Imagem mudou; prepare novamente.')
            with Image.open(source) as image:
                if image.width*image.height>40000000:raise ValueError('Imagem acima do limite de pixels.')
                result=ImageOps.exif_transpose(image).convert('RGB' if path.suffix.lower()=='.jpg' else 'RGBA')
                draw=ImageDraw.Draw(result)
                for mark in marks:
                    if not isinstance(mark,dict) or mark.get('type') not in ('rect','text'):raise ValueError('Anotacao deve ser rect ou text.')
                    color={'cyan':'#45d5d0','yellow':'#f0ce6b','white':'#ffffff'}.get(mark.get('color','cyan'))
                    if color is None:raise ValueError('Cor de anotacao invalida.')
                    coords=mark.get('coords')
                    if not isinstance(coords,list) or len(coords)!=(4 if mark['type']=='rect' else 2) or any(not isinstance(v,(int,float)) or isinstance(v,bool) or not math.isfinite(v) for v in coords):raise ValueError('Coordenadas de anotacao invalidas.')
                    if any(not 0<=v<(result.width if index%2==0 else result.height) for index,v in enumerate(coords)):raise ValueError('Anotacao fora da imagem.')
                    if mark['type']=='rect':
                        if coords[2]<=coords[0] or coords[3]<=coords[1]:raise ValueError('Retangulo invalido.')
                        draw.rectangle(coords,outline=color,width=3)
                    else:draw.text(coords,str(mark.get('text',''))[:500],fill=color)
                if path.suffix.lower()=='.jpg':result=result.convert('RGB')
                result.save(target)
        return request('Anotar imagem' if action=='annotate_image' else 'Converter imagem',str(source)+' → '+str(path),lambda:write_new(path,save))
    raise ValueError('Acao de documentos invalida.')
