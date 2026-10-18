import json
import tempfile
import unittest
from datetime import datetime,timedelta,timezone
from pathlib import Path
from unittest.mock import Mock,patch,AsyncMock
from core import google_services as google,spotify_service as spotify,document_reader as documents,personal_routines as routines

ROOT=Path(__file__).resolve().parents[1]


class IntegrationTests(unittest.TestCase):
    def test_all_six_plugins_discover_without_conflicts(self):
        from core.plugin_loader import discover_plugins
        from core.action_loader import discover_actions
        actions=discover_actions(ROOT/'actions',logger=lambda text:None)
        names=set(actions) if isinstance(actions,dict) else set(actions.names())
        registry=discover_plugins(ROOT/'plugins',names,logger=lambda text:None)
        for name in ('structured_browser','google_calendar','google_drive','spotify_control','document_reader','personal_routines'):
            self.assertTrue(registry.has(name),name)

    def test_tokens_are_encrypted_and_roundtrip(self):
        from core import integration_tokens as tokens
        with tempfile.TemporaryDirectory() as folder,patch.object(tokens,'ROOT',Path(folder)):
            tokens.write('test',{'refresh_token':'ficticio'})
            self.assertNotIn(b'ficticio',(Path(folder)/'token_test.bin').read_bytes())
            self.assertEqual(tokens.read('test')['refresh_token'],'ficticio')
            tokens.remove('test');self.assertEqual(tokens.read('test'),{})

    def test_google_tokens_are_isolated(self):
        with patch.object(google.tokens,'read',return_value={'refresh_token':'ficticio'}) as read,patch.object(google,'settings',return_value={}):
            self.assertTrue(google.status('google_drive')['connected']);read.assert_any_call('google_drive')

    def test_google_wrong_expected_account_blocks_service(self):
        with patch.object(google.tokens,'read',return_value={'jarvis_account_email':'outra@example.test'}),patch.object(google,'settings',return_value={'expected_email':'conta@example.test'}):
            with self.assertRaises(ValueError):google.service('google_calendar')

    def test_google_connect_wrong_account_preserves_existing_token(self):
        credentials=Mock();api=Mock();api.userinfo.return_value.get.return_value.execute.return_value={'email':'outra@example.test'}
        with patch.object(google,'credential_path',return_value=Mock(exists=Mock(return_value=True))),patch.object(google,'client_configuration',return_value={}),patch.object(google,'settings',return_value={'expected_email':'conta@example.test'}),patch('google_auth_oauthlib.flow.InstalledAppFlow.from_client_config') as flow,patch.object(google,'build',return_value=api),patch.object(google.tokens,'write') as save:
            flow.return_value.run_local_server.return_value=credentials
            self.assertFalse(google.connect('google_drive')[0]);save.assert_not_called()
            scopes=flow.call_args.args[1]
            self.assertIn('https://www.googleapis.com/auth/drive.readonly',scopes)
            self.assertNotIn('https://www.googleapis.com/auth/gmail.readonly',scopes)

    def test_google_client_cannot_override_oauth_host(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'client.json'
            path.write_text(json.dumps({'installed':{'client_id':'test.apps.googleusercontent.com','client_secret':'dummy','auth_uri':'https://wrong.example'}}))
            self.assertEqual(google.client_configuration(path)['installed']['auth_uri'],'https://accounts.google.com/o/oauth2/auth')

    def test_calendar_pagination_and_rfc3339(self):
        api=Mock();events=api.events.return_value
        events.list.return_value.execute.return_value={'items':[{'id':'real'}],'nextPageToken':'next'}
        with patch.object(google,'service',return_value=api):
            result=google.calendar({'action':'list','date_from':'2026-10-11T10:00:00-03:00','page_token':'page'})
            self.assertFalse(result['complete']);self.assertEqual(result['next_page_token'],'next')
            self.assertEqual(events.list.call_args.kwargs['pageToken'],'page')
            with self.assertRaises(ValueError):google.calendar({'action':'list','date_from':'2026-10-11T10:00:00'})

    def test_calendar_create_and_update_confirm_real_api(self):
        api=Mock();events=api.events.return_value
        events.insert.return_value.execute.return_value={'id':'real','summary':'Consulta'}
        with patch.object(google,'service',return_value=api):
            result=google.calendar({'action':'create','title':'Consulta','start':'2026-10-12T10:00:00-03:00','end':'2026-10-12T11:00:00-03:00'})
            self.assertTrue(result['confirmed']);self.assertEqual(result['event']['id'],'real')
            with self.assertRaises(ValueError):google.calendar({'action':'create','title':'X','start':'2026-10-12T12:00:00-03:00','end':'2026-10-12T11:00:00-03:00'})
            self.assertEqual(events.insert.call_count,1)
            events.get.return_value.execute.return_value={'start':{'dateTime':'2026-10-12T10:00:00-03:00'},'end':{'dateTime':'2026-10-12T11:00:00-03:00'}}
            google.calendar({'action':'update','event_id':'real','title':'Novo'})
            self.assertEqual(events.patch.call_args.kwargs['body'],{'summary':'Novo'})

    def test_calendar_delete_requires_explicit_confirmation(self):
        api=Mock()
        with patch.object(google,'service',return_value=api):
            with self.assertRaises(ValueError):google.calendar({'action':'delete','event_id':'real'})
        api.events.return_value.delete.assert_not_called()

    def test_drive_query_escape_and_metadata_not_content(self):
        api=Mock();api.files.return_value.list.return_value.execute.return_value={'files':[],'nextPageToken':'more'}
        with patch.object(google,'service',return_value=api):
            result=google.drive({'action':'list','query':"João's"})
            self.assertFalse(result['complete'])
            self.assertIn("João\\'s",api.files.return_value.list.call_args.kwargs['q'])
            api.files.return_value.get.return_value.execute.return_value={'id':'real','name':'Arquivo','mimeType':'text/plain'}
            result=google.drive({'action':'read','file_id':'real'})
            self.assertNotIn('text',result);self.assertIn('Metadados',result['note'])

    def test_drive_download_rejects_large_file_before_request(self):
        api=Mock();api.files.return_value.get.return_value.execute.return_value={'size':str(26*1024*1024)}
        with patch.object(google,'service',return_value=api):
            with self.assertRaises(ValueError):google.drive({'action':'download','file_id':'real'})
        api.files.return_value.get_media.assert_not_called()

    def test_spotify_search_play_uri_and_no_blind_retries(self):
        api=Mock();uri='spotify:track:'+'a'*22
        api.search.return_value={'tracks':{'items':[{'id':'a'*22,'uri':uri,'name':'Música','artists':[{'name':'Artista'}],'album':{'images':[]}}]}}
        api.devices.return_value={'devices':[{'id':'fixture-device','is_active':True}]}
        api.current_playback.return_value={'is_playing':True,'item':{'uri':uri}}
        with patch.object(spotify,'service',return_value=api):
            result=spotify.run({'action':'search','query':'Música'})
            self.assertEqual(result['tracks'][0]['uri'],uri)
            result=spotify.run({'action':'play','uri':uri});self.assertTrue(result['accepted'])
            self.assertNotIn('playing',result)
            with self.assertRaises(ValueError):spotify.run({'action':'play','uri':'inventada'})
            self.assertEqual(api.start_playback.call_count,1)
            with self.assertRaises(ValueError):spotify.run({'action':'volume','volume':101})
            api.volume.assert_not_called()

    def test_spotify_cache_mismatched_app_not_connected(self):
        with patch.object(spotify.tokens,'read',return_value={'refresh_token':'ficticio','jarvis_client_id':'old'}),patch.object(spotify,'config',return_value={'client_id':'new'}):
            self.assertFalse(spotify.status()['connected'])

    def test_spotify_pkce_callback_timeout_and_state(self):
        with patch('spotipy.oauth2.start_local_http_server') as create:
            with patch.object(spotify,'config',return_value={'client_id':'a'*32}):manager=spotify.auth()
            manager._open_auth_url=Mock();server=Mock();server.state='wrong';server.auth_code='test'
            create.return_value.__enter__.return_value=server
            with self.assertRaises(Exception):manager._get_auth_response_local_server(8767)
            self.assertEqual(server.timeout,180)
            server.state=manager.state
            self.assertEqual(manager._get_auth_response_local_server(8767),'test')

    def test_error_does_not_disclose_token(self):
        from core.essential_plugins import respond
        def fail():raise RuntimeError('token-supersecreto')
        value=respond(fail);self.assertNotIn('token-supersecreto',value);self.assertFalse(json.loads(value)['ok'])


class DocumentTests(unittest.TestCase):
    def test_real_portuguese_ocr(self):
        from PIL import Image,ImageDraw,ImageFont
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'imagem.png';image=Image.new('RGB',(1000,200),'white')
            draw=ImageDraw.Draw(image);font=ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf',48)
            draw.text((30,50),'Agenda de amanhã: reunião',font=font,fill='black');image.save(path)
            result=documents.read({'path':str(path)})
            self.assertIn('Agenda',result['text']);self.assertIn('reunião',result['text'])
    def test_read_text_is_paginated(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'texto.txt';path.write_text('a'*500,encoding='utf-8')
            first=documents.read({'path':str(path),'max_chars':100})
            self.assertFalse(first['complete']);self.assertEqual(first['next_offset'],100)
            last=documents.read({'path':str(path),'offset':400,'max_chars':100})
            self.assertTrue(last['complete'])

    def test_docx_xlsx_and_pdf_real_reading(self):
        from docx import Document
        from openpyxl import Workbook
        from PyQt6.QtWidgets import QApplication
        from PyQt6.QtGui import QPdfWriter,QPainter,QFont,QFontDatabase
        app=QApplication.instance() or QApplication([])
        # Offscreen Qt does not discover Windows fonts automatically.
        QFontDatabase.addApplicationFont('C:/Windows/Fonts/segoeui.ttf')
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'arquivo.docx';doc=Document();doc.add_paragraph('Documento de teste');doc.save(path)
            self.assertIn('Documento de teste',documents.read({'path':str(path)})['text'])
            path=Path(folder)/'arquivo.xlsx';book=Workbook();book.active.append(['Matéria','Horas']);book.active.append(['Matemática',3]);book.save(path);book.close()
            self.assertIn('Matemática | 3',documents.read({'path':str(path)})['text'])
            path=Path(folder)/'arquivo.pdf';writer=QPdfWriter(str(path));writer.setResolution(96)
            painter=QPainter(writer);painter.setFont(QFont('Segoe UI',12));painter.drawText(40,80,'Texto PDF verificavel');painter.end()
            del painter,writer
            self.assertIn('Texto PDF verificavel',documents.read({'path':str(path)})['text'])

    def test_download_names_cannot_traverse_or_overwrite(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(documents,'DOWNLOADS',Path(folder)):
            one=documents.save_download('../../teste.txt',b'a');two=documents.save_download('teste.txt',b'b')
            self.assertEqual(one.parent,Path(folder));self.assertNotEqual(one,two)

    def test_gmail_attachment_mismatch_not_downloaded(self):
        from core import gmail_connection as gmail
        api=Mock();api.users.return_value.messages.return_value.get.return_value.execute.return_value={'payload':{'parts':[{'filename':'exemplo.txt','body':{'attachmentId':'real','size':2}}]}}
        with patch.object(gmail,'_service',return_value=api):
            with self.assertRaises(ValueError):documents.gmail_attachment({'message_id':'message','attachment_id':'outra'})
        api.users.return_value.messages.return_value.attachments.assert_not_called()


class RoutineTests(unittest.TestCase):
    def test_schedule_once_is_claimed_and_completed_once(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(routines,'DATABASE',Path(folder)/'test.sqlite3'):
            due=datetime.now(timezone.utc)+timedelta(hours=1)
            result=routines.run({'action':'create','title':'Meu dia','due':due.isoformat(),'domains':['tasks'],'repeat':'once'})
            with patch.object(routines,'briefing',return_value={'ok':True}) as briefing:
                self.assertIsNotNone(routines.poll(due+timedelta(seconds=1)))
                self.assertIsNone(routines.poll(due+timedelta(seconds=2)));briefing.assert_called_once()
            entry=routines.run({'action':'list'})['routines'][0]
            self.assertEqual(entry['id'],result['id']);self.assertFalse(entry['enabled'])

    def test_daily_skips_backlog_and_disable_blocks(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(routines,'DATABASE',Path(folder)/'test.sqlite3'):
            due=datetime.now(timezone.utc)+timedelta(hours=1)
            result=routines.run({'action':'create','title':'Diária','due':due.isoformat(),'repeat':'daily','domains':['gmail']})
            with patch.object(routines,'briefing',return_value={'ok':True}):routines.poll(due+timedelta(days=3))
            entry=routines.run({'action':'list'})['routines'][0]
            self.assertGreater(datetime.fromisoformat(entry['due']),due+timedelta(days=3))
            routines.run({'action':'disable','id':result['id']})
            self.assertIsNone(routines.poll(due+timedelta(days=10)))

    def test_no_arbitrary_actions_or_naive_dates(self):
        with self.assertRaises(ValueError):routines.domains(['send_message'])
        with tempfile.TemporaryDirectory() as folder,patch.object(routines,'DATABASE',Path(folder)/'test.sqlite3'):
            with self.assertRaises(ValueError):routines.run({'action':'create','title':'Teste','due':'2030-01-01T10:00:00'})

    def test_partial_failure_preserves_other_sources(self):
        with patch('actions.organizer.organizer_action',return_value=json.dumps({'ok':True,'data':{'items':[]}})),patch('core.gmail_connection.list_messages',side_effect=RuntimeError('offline')):
            result=routines.briefing(['tasks','gmail'])
            self.assertTrue(result['partial']);self.assertTrue(result['domains']['tasks']['ok'])
            self.assertFalse(result['domains']['gmail']['ok'])


class BrowserTests(unittest.IsolatedAsyncioTestCase):
    async def test_observed_control_real_browser_and_stale_target(self):
        from playwright.async_api import async_playwright
        from actions.browser_control import _BrowserSession
        with patch('actions.browser_control._resolve_browser'):
            session=_BrowserSession('edge',background=True)
        async with async_playwright() as playwright:
            browser=await playwright.chromium.launch(headless=True,channel='msedge')
            page=await browser.new_page()
            await page.set_content('<input aria-label="Buscar"><input type="password"><button onclick="this.textContent=\'Feito\'">Enviar</button>')
            session._get_page=AsyncMock(return_value=page)
            result=await session.observed_controls();self.assertEqual(len(result['controls']),2)
            field=next(item for item in result['controls'] if item['tag']=='INPUT')
            value=await session.observed_action({'action':'fill','control_id':field['control_id'],'text':'Música'})
            self.assertTrue(value['performed']);self.assertEqual(await page.locator('input').first.input_value(),'Música')
            self.assertFalse((await session.observed_action({'action':'fill','control_id':field['control_id'],'text':'outra'}))['ok'])
            result=await session.observed_controls();button=next(item for item in result['controls'] if item['tag']=='BUTTON')
            await page.locator('button').evaluate("e=>e.textContent='Alterado'")
            self.assertFalse((await session.observed_action({'action':'click','control_id':button['control_id']}))['ok'])
            await browser.close()


if __name__=='__main__':unittest.main()
