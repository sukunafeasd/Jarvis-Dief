import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from core import office_resources,outlook_service,notion_service


class CompletionTests(unittest.TestCase):
    def test_audio_transcription_requires_confirmation_and_is_reviewable(self):
        from core import local_resources,resource_store,gemini
        with tempfile.TemporaryDirectory() as folder,patch.object(resource_store,'DATABASE',Path(folder)/'records.sqlite3'):
            audio=Path(folder)/'fixture.wav';audio.write_bytes(b'fixture-not-real-audio')
            reply=Mock();reply.text='Transcricao simulada; sem participantes reais.'
            with patch.object(local_resources,'request',return_value={'completed':False,'job_id':'pending'}),patch.object(gemini,'call') as call:
                result=local_resources.run({'kind':'meeting','action':'transcribe','path':str(audio)})
                self.assertFalse(result['completed']);call.assert_not_called()
            with patch.object(local_resources,'request',side_effect=lambda title,detail,callback:callback()),patch.object(gemini,'call',return_value=reply):
                result=local_resources.run({'kind':'meeting','action':'transcribe','path':str(audio)})
                self.assertTrue(result['review_required']);self.assertFalse(result['item']['metadata']['tasks_created'])
    def test_chart_and_cell_edit_preserve_original_and_literal_formulas(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder,patch.object(office_resources,'request',side_effect=lambda title,detail,callback:callback()):
            source=Path(folder)/'data.xlsx';dest=Path(folder)/'edited.xlsx'
            office_resources.run({'action':'spreadsheet','path':str(source),'rows_json':'[["Categoria","Valor"],["A",-2],["B",4]]','chart':True})
            before=source.read_bytes()
            book=load_workbook(source);self.assertEqual(len(book.active._charts),1);self.assertIsNone(book.active._charts[0].y_axis.scaling.min);book.close()
            office_resources.run({'action':'edit_spreadsheet','other':str(source),'path':str(dest),'edits_json':'[{"cell":"A2","value":"=evil"}]'})
            self.assertEqual(before,source.read_bytes())
            book=load_workbook(dest);self.assertEqual(book.active['A2'].data_type,'s');book.close()
            with self.assertRaises(ValueError):office_resources.run({'action':'edit_spreadsheet','other':str(source),'path':str(Path(folder)/'bad.xlsx'),'edits_json':'[{"cell":"A0","value":4}]'})

    def test_image_annotation_bounds_and_original_preserved(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder,patch.object(office_resources,'request',side_effect=lambda title,detail,callback:callback()):
            source=Path(folder)/'source.png';dest=Path(folder)/'marked.png';Image.new('RGB',(100,100),'black').save(source)
            before=source.read_bytes()
            office_resources.run({'action':'annotate_image','other':str(source),'path':str(dest),'annotations_json':'[{"type":"rect","coords":[10,10,80,80],"color":"cyan"}]'})
            self.assertEqual(source.read_bytes(),before)
            with Image.open(dest) as image:self.assertNotEqual(image.getpixel((10,10))[:3],(0,0,0))
            with self.assertRaises(ValueError):office_resources.run({'action':'annotate_image','other':str(source),'path':str(Path(folder)/'bad.png'),'annotations_json':'[{"type":"rect","coords":[-1,0,20,20]}]'})

    def test_outlook_send_is_not_executed_before_confirmation(self):
        prepared={'account':'fixture@example.com','client_id':'fixture'}
        with patch.object(outlook_service.tokens,'read',return_value=prepared),patch.object(outlook_service,'status',return_value={'connected':True}),patch('core.resource_jobs.request',return_value={'job_id':'pending','completed':False}) as gate,patch.object(outlook_service.requests,'post') as post:
            result=outlook_service.run({'action':'send','to':'recipient@example.com','subject':'Teste','text':'Fixture'})
            self.assertFalse(result['completed']);post.assert_not_called();gate.assert_called_once()

    def test_outlook_send_accepted_is_not_delivery(self):
        prepared={'account':'fixture@example.com','client_id':'fixture'};response=Mock();response.status_code=202
        with patch.object(outlook_service.tokens,'read',return_value=prepared),patch.object(outlook_service,'status',return_value={'connected':True}),patch.object(outlook_service,'access_token',return_value='fixture'),patch('core.resource_jobs.request',side_effect=lambda title,detail,callback:callback()),patch.object(outlook_service.requests,'post',return_value=response):
            result=outlook_service.run({'action':'send','to':'recipient@example.com','subject':'Teste','text':'Fixture'})
        self.assertTrue(result['provider_accepted']);self.assertFalse(result['delivery_confirmed'])

    def test_notion_wrong_identity_preserves_previous_connection(self):
        with patch.object(notion_service.tokens,'read',return_value={'token':'fixture'}),patch.object(notion_service,'call',return_value={'id':'wrong'}),patch.object(notion_service,'settings',return_value={'expected_identity_id':'expected'}),patch.object(notion_service.tokens,'write') as write:
            with self.assertRaises(ValueError):notion_service.connect()
            write.assert_not_called()

    def test_notion_nested_content_not_claimed_complete(self):
        item={'id':'block','type':'paragraph','paragraph':{'rich_text':[{'plain_text':'Fixture'}]},'has_children':True}
        with patch.object(notion_service,'status',return_value={'connected':True}),patch.object(notion_service,'call',return_value={'results':[item],'has_more':False}):
            result=notion_service.run({'action':'read','id':'123456781234123412341234567890ab'})
        self.assertFalse(result['complete']);self.assertTrue(result['items'][0]['has_children'])

    def test_download_limits_and_no_execution(self):
        from core import resource_downloads as downloads,resource_store,document_reader
        with tempfile.TemporaryDirectory() as folder,patch.object(resource_store,'DATABASE',Path(folder)/'records.sqlite3'),patch.object(document_reader,'DOWNLOADS',Path(folder)/'downloads'),patch.object(downloads,'validate_url',side_effect=lambda value:value),patch.object(downloads.requests,'get') as get:
            response=get.return_value.__enter__.return_value;response.status_code=200;response.headers={'Content-Length':'4'};response.iter_content.return_value=[b'data']
            result=downloads.execute('https://example.com/file.txt','file.txt')
            self.assertFalse(result['executed']);self.assertEqual(Path(result['path']).read_bytes(),b'data')
            response.headers={'Content-Length':str(document_reader.MAX_BYTES+1)}
            with self.assertRaises(ValueError):downloads.execute('https://example.com/huge','huge')

    def test_download_refuses_private_destinations_and_redirects(self):
        from core import resource_downloads as downloads
        with patch.object(downloads.socket,'getaddrinfo',return_value=[(2,1,6,'',('127.0.0.1',443))]):
            with self.assertRaises(ValueError):downloads.validate_url('https://internal.example/file')
        with patch.object(downloads,'validate_url',side_effect=lambda value:value),patch.object(downloads.requests,'get') as get:
            get.return_value.__enter__.return_value.status_code=302
            with self.assertRaises(ValueError):downloads.execute('https://example.com/redirect','file')
